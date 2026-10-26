"""
Build DSA_Documentation.docx from DSA_Documentation.md.

Supported Markdown:
  - headings (#, ##, ###, ####)
  - paragraphs, bold, italic, inline code, links
  - bullet and numbered lists
  - pipe tables
  - fenced code blocks
  - images with captions
  - <!-- pagebreak --> and <!-- source-code --> markers
"""

from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

DOCS = Path(__file__).resolve().parent
ROOT = DOCS.parent
MD_PATH = DOCS / "DSA_Documentation.md"
OUT_PATH = DOCS / "DSA_Documentation.docx"
SOURCE_PATH = ROOT / "student_management.cpp"
DOWNLOADS = Path.home() / "Downloads"
DOWNLOADS_COPY = DOWNLOADS / "Data Structure Algorithms - Student Record Management Documentation (REVISED).docx"

NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = RGBColor(0x2E, 0x75, 0xB6)
DARK = RGBColor(0x26, 0x26, 0x26)
GRAY = RGBColor(0x59, 0x59, 0x59)
CODE_RED = RGBColor(0xC7, 0x25, 0x4E)

PAGE_WIDTH_CM = 16.6


# ---------------------------------------------------------------------------
# Low-level XML helpers
# ---------------------------------------------------------------------------

def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def shade_run(run, fill: str) -> None:
    r_pr = run._r.get_or_add_rPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    r_pr.append(shd)


def set_cell_margins(table, top=60, left=90, bottom=60, right=90) -> None:
    tbl_pr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, value in (("top", top), ("left", left),
                        ("bottom", bottom), ("right", right)):
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")
        mar.append(element)
    tbl_pr.append(mar)


def set_table_borders(table, color="D9D9D9", size="4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)
        borders.append(element)
    tbl_pr.append(borders)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def preserve_space(run) -> None:
    run._r.set(qn("xml:space"), "preserve")


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(text)
    run._r.append(end)
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    run.font.color.rgb = GRAY


def shade_paragraph(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def bottom_border(paragraph, color="1F3864", size="8") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


# ---------------------------------------------------------------------------
# Document setup
# ---------------------------------------------------------------------------

def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = DARK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    for name, size, color, before, after in (
        ("Heading 1", 18, NAVY, 18, 8),
        ("Heading 2", 14, BLUE, 14, 5),
        ("Heading 3", 12, NAVY, 10, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri Light"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri Light")


def configure_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(0.9)
    section.footer_distance = Cm(0.9)

    header = section.header
    header_para = header.paragraphs[0]
    header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header_para.add_run(
        "PLMun Student Record Management System  |  DSA Documentation"
    )
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    run.font.color.rgb = GRAY
    bottom_border(header_para, color="BFBFBF", size="4")

    footer = section.footer
    footer_para = footer.paragraphs[0]
    footer_para.paragraph_format.tab_stops.add_tab_stop(
        Cm(PAGE_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT
    )
    left = footer_para.add_run("Group 7  |  BSIT 2H")
    left.font.name = "Calibri"
    left.font.size = Pt(8)
    left.font.color.rgb = GRAY
    footer_para.add_run("\t")
    page_label = footer_para.add_run("Page ")
    page_label.font.name = "Calibri"
    page_label.font.size = Pt(8)
    page_label.font.color.rgb = GRAY
    add_page_number(footer_para)


def add_cover(doc: Document) -> None:
    for _ in range(3):
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Data Structure Algorithms")
    run.font.name = "Calibri Light"
    run.font.size = Pt(30)
    run.font.bold = True
    run.font.color.rgb = NAVY

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Student Record Management Documentation")
    run.font.name = "Calibri Light"
    run.font.size = Pt(20)
    run.font.color.rgb = BLUE

    rule = doc.add_paragraph()
    bottom_border(rule, color="2E75B6", size="12")

    tagline = doc.add_paragraph()
    tagline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = tagline.add_run(
        "A C++17 console application using a singly linked list, stack, queue,\n"
        "linear and binary search, bubble and selection sort, and recursion"
    )
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.italic = True
    run.font.color.rgb = GRAY

    doc.add_paragraph()

    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="D9D9D9")
    set_cell_margins(table, top=90, bottom=90, left=120, right=120)
    rows = [
        ("Group", "Group 7"),
        ("Course / Section", "BSIT 2H"),
        ("Project", "PLMun Student Record Management System"),
        ("Leader", "______________________________"),
        ("Members", "Henry Jr. Banaag\nJanelle Maitim Mansumayog\n"
                    "Emmanuel Jacinto M. Medina\nJetlance Manguiob\n"
                    "Yuan Gabrielle Gatchalian"),
    ]
    for row, (label, value) in zip(table.rows, rows):
        shade_cell(row.cells[0], "DEEAF6")
        p0 = row.cells[0].paragraphs[0]
        r0 = p0.add_run(label)
        r0.font.bold = True
        r0.font.size = Pt(10.5)
        r0.font.color.rgb = NAVY
        p1 = row.cells[1].paragraphs[0]
        lines = value.split("\n")
        for i, line in enumerate(lines):
            if i:
                p1.add_run().add_break()
            r1 = p1.add_run(line)
            r1.font.size = Pt(10.5)
        row.cells[0].width = Cm(3.8)
        row.cells[1].width = Cm(11.8)

    doc.add_paragraph()
    date = doc.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = date.add_run("October 2026")
    run.font.size = Pt(11)
    run.font.color.rgb = GRAY

    doc.add_page_break()


# ---------------------------------------------------------------------------
# Inline formatting
# ---------------------------------------------------------------------------

TOKEN_RE = re.compile(r"(\*\*.+?\*\*|__.+?__|\*[^*\n]+?\*|_[^_\n]+?_|`[^`\n]+?`)")
LINK_RE = re.compile(r"\[([^\]]+)\]\([^)]+\)")


def add_inline_runs(paragraph, text: str, size: float | None = None,
                    color: RGBColor | None = None,
                    bold_all: bool = False) -> None:
    text = LINK_RE.sub(r"\1", text)
    tokens = TOKEN_RE.split(text)
    for token in tokens:
        if not token:
            continue
        bold = bold_all
        italic = False
        code = False
        content = token

        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            content = token[2:-2]
            bold = True
        elif token.startswith("__") and token.endswith("__") and len(token) > 4:
            content = token[2:-2]
            bold = True
        elif token.startswith("*") and token.endswith("*") and len(token) > 3:
            content = token[1:-1]
            italic = True
        elif token.startswith("_") and token.endswith("_") and len(token) > 3:
            content = token[1:-1]
            italic = True
        elif token.startswith("`") and token.endswith("`") and len(token) > 2:
            content = token[1:-1]
            code = True

        run = paragraph.add_run(content)
        run.bold = bold
        run.italic = italic
        if size is not None:
            run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color
        if code:
            run.font.name = "Consolas"
            run.font.size = Pt((size or 10.5) - 1.2)
            run.font.color.rgb = CODE_RED
            shade_run(run, "F2F2F2")
            run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
        if content.startswith(" ") or content.endswith(" "):
            preserve_space(run)


# ---------------------------------------------------------------------------
# Blocks
# ---------------------------------------------------------------------------

def add_code_block(doc: Document, lines: list[str], appendix: bool = False) -> None:
    max_len = max((len(line) for line in lines), default=0)
    if appendix:
        font_size = 7.0
        line_spacing = 0.98
    elif len(lines) > 80 or max_len > 100:
        font_size = 7.6
        line_spacing = 1.0
    else:
        font_size = 8.6
        line_spacing = 1.0

    if appendix:
        # A single-cell table with thousands of paragraphs is slow in Word and
        # may not break cleanly across pages. Shaded body paragraphs are more
        # reliable for the full source listing.
        for line in lines:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = line_spacing
            paragraph.paragraph_format.left_indent = Cm(0.15)
            paragraph.paragraph_format.right_indent = Cm(0.15)
            shade_paragraph(paragraph, "F7F7F7")
            run = paragraph.add_run(line if line else " ")
            run.font.name = "Consolas"
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)
            run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
            preserve_space(run)
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(4)
        spacer.paragraph_format.space_before = Pt(0)
        return

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color="D9D9D9", size="4")
    set_cell_margins(table, top=80, left=100, bottom=80, right=100)

    cell = table.cell(0, 0)
    cell.width = Cm(PAGE_WIDTH_CM)
    shade_cell(cell, "F7F7F7")

    first = cell.paragraphs[0]
    first._p.getparent().remove(first._p)

    for line in lines:
        paragraph = cell.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = line_spacing
        run = paragraph.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(font_size)
        run.font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
        preserve_space(run)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)
    spacer.paragraph_format.space_before = Pt(0)


def add_image(doc: Document, alt: str, path: Path) -> None:
    if not path.exists():
        return
    name = path.name
    width = Inches(6.3)
    if "flowchart" in name:
        width = Inches(5.3)
    elif "linked_list" in name:
        width = Inches(6.4)
    elif name == "stack.png":
        width = Inches(3.9)
    elif name == "queue.png":
        width = Inches(6.2)
    elif "complexity" in name:
        width = Inches(6.4)

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run().add_picture(str(path), width=width)

    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(10)
    run = caption.add_run(alt)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = GRAY


def add_table(doc: Document, raw_rows: list[list[str]]) -> None:
    if not raw_rows:
        return
    cols = max(len(row) for row in raw_rows)
    rows = [row + [""] * (cols - len(row)) for row in raw_rows]

    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_cell_margins(table, top=55, left=80, bottom=55, right=80)

    font_size = 9.0 if cols <= 4 else 8.2 if cols <= 6 else 7.6

    # Heuristic column widths based on the longest cell in each column.
    lengths = []
    for c in range(cols):
        lengths.append(max(6, min(42, max(len(row[c]) for row in rows))))
    total = sum(lengths)
    widths = [max(1.15, PAGE_WIDTH_CM * length / total) for length in lengths]
    scale = PAGE_WIDTH_CM / sum(widths)
    widths = [w * scale for w in widths]

    for r_index, row in enumerate(rows):
        for c_index, text in enumerate(row):
            cell = table.cell(r_index, c_index)
            cell.width = Cm(widths[c_index])
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            if r_index == 0:
                shade_cell(cell, "1F3864")
                add_inline_runs(paragraph, text, size=font_size,
                                color=RGBColor(0xFF, 0xFF, 0xFF),
                                bold_all=True)
            else:
                if r_index % 2 == 0:
                    shade_cell(cell, "F2F2F2")
                add_inline_runs(paragraph, text, size=font_size)
        if r_index == 0:
            repeat_header(table.rows[0])

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)
    spacer.paragraph_format.space_before = Pt(0)


def add_list_item(doc: Document, indent: int, marker: str, text: str) -> None:
    paragraph = doc.add_paragraph()
    level = min(indent // 2, 2)
    paragraph.paragraph_format.left_indent = Cm(0.55 + 0.55 * level)
    paragraph.paragraph_format.first_line_indent = Cm(-0.42)
    paragraph.paragraph_format.space_after = Pt(2.5)
    marker_run = paragraph.add_run(marker + "  ")
    marker_run.bold = True
    marker_run.font.color.rgb = BLUE
    add_inline_runs(paragraph, text, size=10.5)


# ---------------------------------------------------------------------------
# Main Markdown parser
# ---------------------------------------------------------------------------

def build() -> None:
    if not MD_PATH.exists():
        raise SystemExit(f"Missing {MD_PATH}")

    text = MD_PATH.read_text(encoding="utf-8")
    lines = text.splitlines()

    # Skip the cover block (everything before the Table of Contents).
    start = 0
    for i, line in enumerate(lines):
        if line.strip().lower() == "## table of contents":
            start = i
            break
    lines = lines[start:]

    doc = Document()
    configure_styles(doc)
    configure_page(doc)
    add_cover(doc)

    doc.core_properties.title = (
        "Data Structure Algorithms - Student Record Management Documentation"
    )
    doc.core_properties.author = "Group 7 - BSIT 2H"
    doc.core_properties.subject = "Data Structures and Algorithms"
    doc.core_properties.comments = (
        "Revised documentation for the PLMun Student Record Management System."
    )

    i = 0
    total = len(lines)
    first_section = True

    while i < total:
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped.startswith("<!--"):
            if "pagebreak" in stripped:
                doc.add_page_break()
            elif "source-code" in stripped:
                if SOURCE_PATH.exists():
                    source_lines = SOURCE_PATH.read_text(encoding="utf-8").splitlines()
                    add_code_block(doc, source_lines, appendix=True)
                    doc.add_page_break()
                i += 1
                continue
            i += 1
            continue

        if stripped.startswith("```"):
            language = stripped[3:].strip()
            code_lines: list[str] = []
            i += 1
            while i < total and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1
            add_code_block(doc, code_lines)
            continue

        image_match = re.fullmatch(r"!\[(.*?)\]\((.*?)\)", stripped)
        if image_match:
            alt, relative = image_match.groups()
            add_image(doc, alt, DOCS / relative)
            i += 1
            continue

        if stripped.startswith("|"):
            table_rows: list[list[str]] = []
            while i < total and lines[i].strip().startswith("|"):
                row_line = lines[i].strip()
                cells = [cell.strip() for cell in row_line.strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", cell) for cell in cells):
                    table_rows.append(cells)
                i += 1
            add_table(doc, table_rows)
            continue

        heading = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if heading:
            level = len(heading.group(1))
            title = heading.group(2).strip()
            if level == 1:
                if not first_section:
                    doc.add_page_break()
                first_section = False
                doc.add_heading(title, level=1)
            elif level == 2:
                is_main = bool(re.match(r"^\d+\.", title)) or \
                    title.lower() == "table of contents"
                if title.startswith("1. Introduction") or title.startswith("15. Appendices"):
                    doc.add_page_break()
                doc.add_heading(title, level=1 if is_main else 2)
            elif level == 3:
                doc.add_heading(title, level=2)
            else:
                doc.add_heading(title, level=3)
            i += 1
            continue

        bullet = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        if bullet:
            add_list_item(doc, len(bullet.group(1)), "•", bullet.group(2))
            i += 1
            continue

        numbered = re.match(r"^(\s*)(\d+\.)\s+(.*)$", line)
        if numbered:
            add_list_item(doc, len(numbered.group(1)), numbered.group(2),
                          numbered.group(3))
            i += 1
            continue

        if stripped.startswith(">"):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.6)
            paragraph.paragraph_format.space_after = Pt(6)
            run = paragraph.add_run(stripped.lstrip("> ").strip())
            run.italic = True
            run.font.color.rgb = GRAY
            i += 1
            continue

        if stripped == "---":
            i += 1
            continue

        # Paragraph: collect consecutive plain lines.
        paragraph_lines = [stripped]
        i += 1
        while i < total:
            nxt = lines[i]
            nxt_stripped = nxt.strip()
            if (not nxt_stripped or nxt_stripped.startswith(("<!--", "```", "|",
                                                             "#", ">", "---"))
                    or re.match(r"^\s*([-*]|\d+\.)\s+", nxt)):
                break
            paragraph_lines.append(nxt_stripped)
            i += 1

        paragraph = doc.add_paragraph()
        add_inline_runs(paragraph, " ".join(paragraph_lines))

    doc.save(OUT_PATH)
    print(f"Wrote {OUT_PATH}")

    if DOWNLOADS.exists():
        shutil.copy2(OUT_PATH, DOWNLOADS_COPY)
        print(f"Copied to {DOWNLOADS_COPY}")

    print(f"Paragraphs: {len(doc.paragraphs)}")
    print(f"Tables: {len(doc.tables)}")
    print(f"Size: {OUT_PATH.stat().st_size / 1024:.0f} KB")


if __name__ == "__main__":
    build()
